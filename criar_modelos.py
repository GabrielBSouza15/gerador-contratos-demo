"""Modelos inteiramente fictícios para demonstrar substituição de campos no Word."""
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent


def criar_modelo(pagamento, destino):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for nome in ['Normal', 'Title', 'Heading 1']:
        estilo = doc.styles[nome]
        estilo.font.name = 'Arial'
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.font.size = Pt(10.5 if nome == 'Normal' else 12 if nome == 'Heading 1' else 17)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.line_spacing = 1.05
    doc.add_paragraph('Contrato demonstrativo de serviços digitais', 'Title')
    doc.add_paragraph('Exemplo fictício destinado à demonstração de preenchimento automático. Não representa contratação real nem modelo jurídico para uso profissional.')
    doc.add_paragraph('Registro -id- | Modalidade ' + ('cartão' if pagamento == 1 else 'PIX'))
    doc.add_paragraph('Prestadora fictícia: Oficina de Dados Exemplo, CNPJ 00.000.000/0000-00, Rua de Demonstração, 100, Cidade Exemplo, UF XX, CEP 00000-000. Contato: contato@example.com.')
    doc.add_paragraph('Solicitante fictício: -qualificacao-.')
    conteudo = [
        ('Atividade de demonstração', 'O cenário considera a atividade -tipo-. O resultado de exemplo reúne um arquivo de treino e um roteiro curto para explorar seus recursos com dados inventados.'),
        ('Agenda e apresentação', 'A demonstração prevê -prazo- dias corridos a partir do início combinado no exercício. Uma apresentação de até quarenta minutos utiliza exclusivamente os materiais de treino.'),
        ('Valores do exercício', 'O valor ilustrativo é R$ -valor- (-extenso-). ' +
         ('O cenário representa pagamento único por cartão, sem conexão com operadora ou coleta de dados do cartão.' if pagamento == 1 else
          'O exercício divide o valor em entrada de R$ -entrada- (-extensoentrada-) e saldo de R$ -saldo- (-extensosaldo-). O saldo corresponde à apresentação final. Não há chave PIX ou cobrança real.')),
        ('Conferência do exemplo', 'A conferência consiste em abrir o arquivo de treino e percorrer três situações fictícias descritas no roteiro. Uma rodada de comentários pode indicar até dois ajustes de apresentação dentro desse exercício.'),
        ('Materiais e dados utilizados', 'Todos os registros do cenário são sintéticos. O material não inclui documentos de clientes, acessos a sistemas empresariais ou arquivos operacionais de terceiros.'),
        ('Conclusão da demonstração', 'O exercício termina com a apresentação dos arquivos de treino. Um novo assunto ou uma nova base de exemplo exige outro registro demonstrativo. Este documento não autoriza assinatura, transferência financeira ou prestação de serviço real.'),
    ]
    for titulo, texto in conteudo:
        doc.add_paragraph(titulo, 'Heading 1')
        doc.add_paragraph(texto)
    doc.add_paragraph('Data do documento demonstrativo: -dia- de -mes- de -ano-.')
    for parte in (doc._element, doc.styles.element):
        for borda in list(parte.iter(qn('w:pBdr'))):
            borda.getparent().remove(borda)
    props = doc.core_properties
    props.author = props.last_modified_by = ''
    props.title = 'Contrato demonstrativo de serviços digitais'
    props.subject = 'Exemplo fictício de automação documental'
    props.comments = ''
    props.created = props.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)


if __name__ == '__main__':
    criar_modelo(1, BASE / 'modelos/contrato_cartao.docx')
    criar_modelo(2, BASE / 'modelos/contrato_pix.docx')
