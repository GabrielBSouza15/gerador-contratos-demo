"""Geração local de um contrato demonstrativo a partir de um ID da planilha."""
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re
import sys

from docx import Document
from num2words import num2words
import pandas as pd

BASE = Path(__file__).resolve().parent
PLANILHA = BASE / 'dados' / 'Informacoes.xlsx'
MESES = dict(enumerate(['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho',
                       'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro'], 1))
SERVICOS = {1: 'Organização de planilhas de exemplo em Excel',
            2: 'Construção de um relatório demonstrativo em Power BI',
            3: 'Criação de uma tela demonstrativa em Power Apps'}
CAMPOS = ['ID', 'Responsável', 'Nome completo', 'CPF', 'Nacionalidade', 'Estado civil',
          'Rua1', 'Complemento1', 'Bairro1', 'Cidade1', 'Estado1', 'CEP1',
          'Razão social', 'CNPJ', 'Rua', 'Complemento', 'Bairro', 'Cidade', 'Estado',
          'CEP', 'Valor', 'Prazo']


def formatar_cnpj(cnpj):
    cnpj = re.sub(r'\D', '', str(cnpj))
    return f'{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}' if len(cnpj) == 14 else cnpj


def formatar_cpf(cpf):
    cpf = re.sub(r'\D', '', str(cpf))
    return f'{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}' if len(cpf) == 11 else cpf


def formatar_cep(cep):
    cep = re.sub(r'\D', '', str(cep))
    return f'{cep[:5]}-{cep[5:]}' if len(cep) == 8 else cep


def valor_decimal(valor):
    try:
        numero = Decimal(str(valor).strip())
        if not numero.is_finite() or numero <= 0 or numero > Decimal('999999999.99'):
            raise ValueError
        if numero != numero.quantize(Decimal('0.01')):
            raise ValueError
        return numero.quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        raise ValueError('Valor deve ser positivo, com no máximo duas casas decimais. Use uma célula numérica ou ponto decimal.') from None


def formatar_valor(valor):
    return f'{Decimal(str(valor)):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def valor_por_extenso(valor):
    return num2words(Decimal(str(valor)), lang='pt_BR', to='currency')


def validar_cpf(cpf):
    cpf = re.sub(r'\D', '', str(cpf))
    if len(cpf) != 11 or cpf == cpf[0] * 11:
        return False
    return all((sum(int(cpf[j]) * (i + 1 - j) for j in range(i)) * 10 % 11) % 10 == int(cpf[i])
               for i in range(9, 11))


def validar_cnpj(cnpj):
    cnpj = re.sub(r'\D', '', str(cnpj))
    if len(cnpj) != 14 or cnpj == cnpj[0] * 14:
        return False
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        resto = sum(int(cnpj[j]) * pesos[j] for j in range(len(pesos))) % 11
        if (0 if resto < 2 else 11 - resto) != int(cnpj[len(pesos)]):
            return False
    return True


def selecionar_por_id(planilha, identificador):
    identificador = str(identificador).strip().upper()
    if not re.fullmatch(r'CT-[0-9]{3}', identificador):
        raise ValueError('Informe um ID no formato CT-001.')
    # dtype e keep_default_na preservam zeros e campos opcionais vazios.
    tabela = pd.read_excel(planilha, sheet_name='Dados', dtype=str, keep_default_na=False)
    faltantes = set(CAMPOS) - set(tabela.columns)
    if faltantes:
        raise ValueError('Colunas ausentes: ' + ', '.join(sorted(faltantes)))
    selecionadas = tabela.loc[tabela['ID'].str.strip().str.upper() == identificador]
    if len(selecionadas) != 1:
        raise ValueError('ID não encontrado.' if selecionadas.empty else 'ID duplicado. Corrija a planilha antes de gerar.')
    return selecionadas.iloc[0].str.strip().to_dict()


def substituir_em_runs(paragrafo, referencias):
    """Substitui inclusive campos divididos em runs, preservando o texto ao redor."""
    for codigo, valor in referencias.items():
        texto = ''.join(run.text for run in paragrafo.runs)
        ocorrencias = list(re.finditer(re.escape(codigo), texto))
        for ocorrencia in reversed(ocorrencias):
            inicio, fim = ocorrencia.span()
            posicao = 0
            inserido = False
            for run in paragrafo.runs:
                anterior = run.text
                final = posicao + len(anterior)
                if posicao < fim and final > inicio:
                    esquerda = max(0, inicio - posicao)
                    direita = min(len(anterior), fim - posicao)
                    run.text = anterior[:esquerda] + (str(valor) if not inserido else '') + anterior[direita:]
                    inserido = True
                posicao = final


def paragrafos_documento(documento):
    yield from documento.paragraphs
    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                yield from celula.paragraphs


def preparar_referencias(linha, pessoa, servico, hoje):
    obrigatorios = ['Responsável', 'Nome completo', 'CPF', 'Nacionalidade', 'Estado civil',
                    'Rua1', 'Cidade1', 'Estado1', 'CEP1', 'Valor', 'Prazo']
    if pessoa == 1:
        obrigatorios += ['Razão social', 'CNPJ', 'Rua', 'Cidade', 'Estado', 'CEP']
    if any(not linha[campo] for campo in obrigatorios):
        raise ValueError('Preencha os campos obrigatórios para o tipo de pessoa escolhido.')
    if not linha['Prazo'].isdigit() or not 1 <= int(linha['Prazo']) <= 3650:
        raise ValueError('Prazo deve ser um número inteiro de dias entre 1 e 3650.')
    total = valor_decimal(linha['Valor'])
    entrada = (total / 2).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    saldo = total - entrada
    endereco = ', '.join(linha[c] for c in ['Rua1', 'Complemento1', 'Bairro1', 'Cidade1', 'Estado1'] if linha[c])
    pessoa_texto = (f"{linha['Nome completo']}, CPF {formatar_cpf(linha['CPF'])}, "
                    f"nacionalidade {linha['Nacionalidade']}, estado civil {linha['Estado civil']}, "
                    f"endereço {endereco}, CEP {formatar_cep(linha['CEP1'])}")
    if pessoa == 1:
        sede = ', '.join(linha[c] for c in ['Rua', 'Complemento', 'Bairro', 'Cidade', 'Estado'] if linha[c])
        qualificacao = (f"{linha['Razão social']}, CNPJ {formatar_cnpj(linha['CNPJ'])}, "
                         f"sede {sede}, CEP {formatar_cep(linha['CEP'])}, representada por {pessoa_texto}")
    else:
        qualificacao = pessoa_texto
    referencias = {'-id-': linha['ID'].upper(), '-qualificacao-': qualificacao,
        '-tipo-': SERVICOS[servico], '-prazo-': str(int(linha['Prazo'])),
        '-valor-': formatar_valor(total), '-extenso-': valor_por_extenso(total),
        '-entrada-': formatar_valor(entrada), '-saldo-': formatar_valor(saldo),
        '-extensoentrada-': valor_por_extenso(entrada), '-extensosaldo-': valor_por_extenso(saldo),
        '-dia-': str(hoje.day), '-mes-': MESES[hoje.month], '-ano-': str(hoje.year)}
    return referencias


def gerar_contrato(identificador, pagamento, pessoa, servico, *, planilha=PLANILHA, saida=BASE / 'gerados', hoje=None):
    if pagamento not in (1, 2) or pessoa not in (1, 2) or servico not in SERVICOS:
        raise ValueError('Escolha uma opção válida nos menus.')
    linha = selecionar_por_id(planilha, identificador)
    hoje = hoje or datetime.now()
    referencias = preparar_referencias(linha, pessoa, servico, hoje)
    avisos = []
    if not validar_cpf(linha['CPF']):
        avisos.append('CPF não passou na validação matemática; os exemplos usam zeros deliberadamente.')
    if pessoa == 1 and not validar_cnpj(linha['CNPJ']):
        avisos.append('CNPJ não passou na validação matemática; os exemplos usam zeros deliberadamente.')
    modelo = BASE / 'modelos' / ('contrato_cartao.docx' if pagamento == 1 else 'contrato_pix.docx')
    documento = Document(modelo)  # Um modelo novo para a única linha selecionada.
    for paragrafo in paragrafos_documento(documento):
        substituir_em_runs(paragrafo, referencias)
    if re.search(r'-[a-z]+-', '\n'.join(p.text for p in paragrafos_documento(documento))):
        raise ValueError('Há campo não preenchido no modelo. O documento não foi salvo.')
    documento.core_properties.author = ''
    documento.core_properties.last_modified_by = ''
    saida = Path(saida)
    saida.mkdir(parents=True, exist_ok=True)
    nome = f"Contrato_{linha['ID'].upper()}_{hoje:%Y%m%d_%H%M%S_%f}"
    destino = saida / (nome + '.docx')
    with destino.open('xb') as arquivo:
        documento.save(arquivo)
    return destino, avisos


def main():
    print('Gerador de contratos demonstrativos — dados fictícios')
    identificador = input('ID do registro (ex.: CT-001): ')
    try:
        pagamento = int(input('Pagamento: 1 - Cartão; 2 - PIX: '))
        pessoa = int(input('Pessoa: 1 - Jurídica; 2 - Física: '))
        print('Serviços: 1 - Excel; 2 - Power BI; 3 - Power Apps')
        servico = int(input('Selecione o serviço: '))
        caminho, avisos = gerar_contrato(identificador, pagamento, pessoa, servico)
        for aviso in avisos:
            print('Aviso:', aviso)
        print('Documento gerado:', caminho)
    except (ValueError, OSError, KeyError) as erro:
        print('Não foi possível gerar:', erro)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
