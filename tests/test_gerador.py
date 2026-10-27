from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest

import pandas as pd
from docx import Document
import gerar_contrato as g


class GeradorTests(unittest.TestCase):
    def test_ids_select_distinct_rows_and_preserve_zero_documents(self):
        for identifier in ('CT-001','CT-002','CT-003','CT-004'):
            row=g.selecionar_por_id(g.PLANILHA,identifier)
            self.assertEqual(row['ID'],identifier)
            self.assertEqual(g.re.sub(r'\D','',row['CPF']),'00000000000')
        self.assertEqual(g.selecionar_por_id(g.PLANILHA,' ct-002 ')['Razão social'],'')

    def test_unknown_blank_and_malformed_id_fail(self):
        for identifier in ('CT-999','','../../contrato','CT-1'):
            with self.subTest(identifier=identifier),self.assertRaises(ValueError):
                g.selecionar_por_id(g.PLANILHA,identifier)

    def test_duplicate_id_fails_instead_of_first_match(self):
        df=pd.read_excel(g.PLANILHA,sheet_name='Dados',dtype=str,keep_default_na=False)
        duplicated=pd.concat([df.iloc[:1],df.iloc[:1]],ignore_index=True)
        with patch.object(g.pd,'read_excel',return_value=duplicated),self.assertRaisesRegex(ValueError,'duplicado'):
            g.selecionar_por_id('unused.xlsx','CT-001')

    def test_odd_cent_preserves_total_and_exact_extenso(self):
        row=g.selecionar_por_id(g.PLANILHA,'CT-003')
        references=g.preparar_referencias(row,1,1,datetime(2026,10,27))
        self.assertEqual(references['-entrada-'],'1.150,01')
        self.assertEqual(references['-saldo-'],'1.150,00')
        self.assertEqual(g.valor_por_extenso(Decimal('1.01')),'um real e um centavo')
        self.assertIn('outubro',references['-mes-'])

    def test_generation_selects_only_chosen_person_for_all_menu_combinations(self):
        with TemporaryDirectory() as folder:
            for payment in (1,2):
                for person,identifier,expected in ((1,'CT-001','Pessoa Exemplo Um'),(2,'CT-002','Pessoa Exemplo Dois')):
                    for service in (1,2,3):
                        target,warnings=g.gerar_contrato(identifier,payment,person,service,saida=folder)
                        text='\n'.join(p.text for p in Document(target).paragraphs)
                        self.assertIn(expected,text)
                        self.assertIn(g.SERVICOS[service],text)
                        self.assertNotIn('Pessoa Exemplo Três',text)
                        self.assertNotRegex(text,r'-[a-z]+-')
                        self.assertIn('demonstração',text)
                        self.assertTrue(warnings)
            self.assertEqual(len(list(Path(folder).glob('*.docx'))),12)

    def test_missing_fields_and_bad_values_fail_without_output(self):
        row=g.selecionar_por_id(g.PLANILHA,'CT-001')
        for change in ({'Nome completo':''},{'Valor':'NaN'},{'Valor':'-1'},{'Valor':'1.001'},{'Prazo':'1.5'}):
            with self.subTest(change=change),self.assertRaises(ValueError):
                g.preparar_referencias(row|change,1,1,datetime.now())
        with TemporaryDirectory() as folder:
            for args in [('CT-001',9,1,1),('CT-002',1,1,1)]:
                with self.assertRaises(ValueError):g.gerar_contrato(*args,saida=folder)
            self.assertFalse(list(Path(folder).glob('*')))

    def test_same_timestamp_refuses_overwrite(self):
        with TemporaryDirectory() as folder:
            options=dict(saida=folder,hoje=datetime(2026,10,27,10,20,30))
            target,_=g.gerar_contrato('CT-001',2,1,1,**options)
            before=target.read_bytes()
            with self.assertRaises(FileExistsError):g.gerar_contrato('CT-001',2,1,1,**options)
            self.assertEqual(before,target.read_bytes())

    def test_split_run_placeholder_keeps_surrounding_formatting(self):
        doc=Document()
        p=doc.add_paragraph()
        p.add_run('Antes ').bold=True
        p.add_run('-no')
        p.add_run('me-')
        p.add_run(' depois').italic=True
        g.substituir_em_runs(p,{'-nome-':'Exemplo'})
        self.assertEqual(p.text,'Antes Exemplo depois')
        self.assertTrue(p.runs[0].bold)
        self.assertTrue(p.runs[-1].italic)

    def test_interactive_flow_requests_id_and_generates_one_document(self):
        with patch('builtins.input',side_effect=['CT-002','1','2','2']) as prompt,patch.object(g,'gerar_contrato',return_value=(Path('exemplo.docx'),[])) as generate:
            self.assertEqual(g.main(),0)
            self.assertIn('ID',prompt.call_args_list[0].args[0])
            generate.assert_called_once_with('CT-002',1,2,2)


if __name__=='__main__':unittest.main()
